"""厚労省リスト(data/01-05.xlsx)＋医薬品マスター(data/y.zip) → docs/index.html（先発・後発 検索）"""
import csv, io, json, re, unicodedata, zipfile
from collections import defaultdict
from datetime import datetime, timedelta, timezone
from pathlib import Path

import openpyxl

ROOT = Path(__file__).parent
DATA = ROOT / "data"

# フラグ（ビット）
F_STAR, F_HOSHI, F_JUN, F_SENTEI, F_KEIKA, F_NOGE, F_SAME = 1, 2, 4, 8, 16, 32, 64


def nk(s):
    return unicodedata.normalize("NFKC", str(s or "")).strip()


def load_list():
    """薬価基準収載品目リスト（内用・注射・外用・歯科）＋「その他」(後発品の有無 1/2/3/★/☆)"""
    rows = {}
    for n in "1234":
        ws = openpyxl.load_workbook(DATA / f"0{n}.xlsx", read_only=True, data_only=True).worksheets[0]
        it = ws.iter_rows(values_only=True)
        hdr = [str(h or "") for h in next(it)]

        def col(key, default):
            for i, h in enumerate(hdr):
                if key in h:
                    return i
            return default
        c = dict(kubun=col("区分", 0), code=col("薬価基準収載医薬品コード", 1), seibun=col("成分名", 2),
                 kikaku=col("規格", 3), name=col("品名", 7), maker=col("メーカー", 8),
                 kouhatsu=col("算定対象となる後発医薬品", 9), douitsu=col("同一剤形", 11),
                 keika=col("経過措置", 13))
        # 「先発医薬品」単独の列（「同一剤形・規格の後発医薬品がある先発医薬品」と区別）
        c["senpatsu"] = next((i for i, h in enumerate(hdr) if h.strip() == "先発医薬品"), 10)
        for r in it:
            code = r[c["code"]]
            if not code:
                continue
            rows[str(code).strip()] = {k: r[i] for k, i in c.items()}
    umu = {}
    ws = openpyxl.load_workbook(DATA / "05.xlsx", read_only=True, data_only=True).worksheets[0]
    it = ws.iter_rows(values_only=True)
    hdr = [str(h or "") for h in next(it)]
    ci = next((i for i, h in enumerate(hdr) if "後発医薬品の有無" in h), 3)
    for r in it:
        if r[0]:
            umu[str(r[0]).strip()] = str(r[ci]).strip() if r[ci] is not None else ""
    return rows, umu


def load_y():
    z = zipfile.ZipFile(DATA / "y.zip")
    name = [n for n in z.namelist() if n.lower().endswith(".csv")][0]
    rows = list(csv.reader(io.TextIOWrapper(z.open(name), encoding="cp932")))
    m = re.search(r"(\d{8})", name)
    return rows, (m.group(1) if m else "")


def name_key(s):
    return re.sub(r"\s", "", unicodedata.normalize("NFKC", str(s or ""))).upper()


def load_yj():
    p = DATA / "yj_names.json"
    return json.loads(p.read_text(encoding="utf-8")) if p.exists() else {}


def build():
    lst, umu = load_list()
    y, ydate = load_y()
    yjn = load_yj()
    # 有効なレコードのみ（廃止、選定療養用の「（選）」「（類）」レコードを除く）
    act = [r for r in y if r[0] != "9" and r[30] == "99999999" and r[41] != "2"
           and not re.search(r"（[選類]）$", r[4]) and r[31] in lst]
    referenced = {r[22] for r in act if r[22] not in ("", "0")}   # 銘柄から参照される統一名レコード
    seen_codes = set()

    # グループ化: 薬価基準コード先頭9桁（成分・剤形・規格）＋一般名コードで連結
    parent = {}

    def find(x):
        parent.setdefault(x, x)
        while parent[x] != x:
            parent[x] = parent[parent[x]]
            x = parent[x]
        return x

    def union(a, b):
        ra, rb = find(a), find(b)
        if ra != rb:
            parent[max(ra, rb)] = min(ra, rb)

    for r in act:
        yk = r[31]
        find("P" + yk[:9])
        if r[36]:
            union("P" + yk[:9], "I" + r[36])

    groups = defaultdict(lambda: {"items": [], "ippan": set(), "touitsu": set(), "kikaku": set(), "unit": ""})
    for r in act:
        code, yk = r[2], r[31]
        L = lst[yk]
        g = groups[find("P" + yk[:9])]
        if r[37]:
            g["ippan"].add(nk(r[37]).replace("【般】", ""))
        g["kikaku"].add(nk(L["kikaku"]))
        g["unit"] = g["unit"] or nk(r[9])
        if not L["maker"]:
            g["touitsu"].add(nk(L["name"]))
        if code in referenced:          # 統一名レコード本体は銘柄があるので出さない
            continue
        seen_codes.add(yk)
        kou, sen, u = L["kouhatsu"], L["senpatsu"], umu.get(yk, "")
        flags = 0
        if kou in ("後発品", "★"):
            cat = "G"
            if kou == "★":
                flags |= F_STAR
        elif sen in ("先発品", "準先発品"):
            cat = "B"
            if sen == "準先発品":
                flags |= F_JUN
            if u == "☆":
                flags |= F_HOSHI
            if u == "1":
                flags |= F_NOGE
            if L["douitsu"]:
                flags |= F_SAME
        else:
            cat = "O"
        if r[41] == "1":
            flags |= F_SENTEI
        keika = nk(L["keika"])
        if keika:
            flags |= F_KEIKA
        try:
            price = float(r[11])
            price = int(price) if price == int(price) else price
        except ValueError:
            price = None
        # 添付文書リンク用のYJコード: 銘柄別収載は薬価基準コードと同じ(1)。統一名収載の銘柄は品名から引く。不明は0
        if L["maker"]:
            yj = 1
        else:
            c = [x for x in yjn.get(name_key(r[34] or r[4]), "").split(",") if x[:9] == yk[:9]]
            yj = max(c) if c else 0
        g["items"].append([nk(r[34] or r[4]), nk(L["maker"]), price, cat, flags, nk(r[6]), yk, keika, yj])

    # 成分（薬価基準コード先頭7桁）ごとにまとめる
    ings = defaultdict(lambda: {"seibun": "", "kubun": "", "groups": []})
    order = {"B": 0, "G": 1, "O": 2}
    for key, g in groups.items():
        if not g["items"]:
            continue
        g["items"].sort(key=lambda i: (order[i[3]], bool(i[4] & F_STAR), i[0]))
        yk = g["items"][0][6]
        L = lst[yk]
        ing = ings[yk[:7]]
        ing["seibun"] = ing["seibun"] or nk(L["seibun"]) or g["items"][0][0]
        ing["kubun"] = ing["kubun"] or nk(L["kubun"])
        ing["groups"].append([min(i[6] for i in g["items"])[:9], " / ".join(sorted(g["kikaku"])),
                              " / ".join(sorted(g["ippan"])) or " / ".join(sorted(g["touitsu"])),
                              g["unit"], g["items"]])
    out = []
    for k7, ing in ings.items():
        ing["groups"].sort(key=lambda g: g[0])
        out.append([k7, ing["seibun"], ing["kubun"], [g[1:] for g in ing["groups"]]])
    out.sort(key=lambda i: (i[1], i[0]))

    src = {}
    sp = DATA / "source.json"
    if sp.exists():
        src = json.loads(sp.read_text(encoding="utf-8"))
    n_items = sum(len(g[3]) for i in out for g in i[3])
    meta = {"applied": nk(src.get("applied", "")), "fetched": src.get("fetched", ""), "page": src.get("page", ""),
            "ydate": ydate, "items": n_items, "ings": len(out),
            "doc": sum(1 for i in out for g in i[3] for it in g[3] if it[8]),
            "brand": sum(1 for i in out for g in i[3] for it in g[3] if it[3] == "B"),
            "generic": sum(1 for i in out for g in i[3] for it in g[3] if it[3] == "G")}
    return {"meta": meta, "ings": out}


def main():
    d = build()
    html = (ROOT / "template.html").read_text(encoding="utf-8")
    js = json.dumps(d, ensure_ascii=False, separators=(",", ":")).replace("</", "<\\/")
    (ROOT / "docs").mkdir(exist_ok=True)
    (ROOT / "docs" / "index.html").write_text(html.replace("__DATA__", js), encoding="utf-8")
    # 自動更新が最後に動いた時刻（データに変更がない日も更新される）
    now = datetime.now(timezone(timedelta(hours=9))).strftime("%Y-%m-%d %H:%M")
    (ROOT / "docs" / "status.json").write_text(json.dumps({"checked": now}), encoding="utf-8")
    print(d["meta"])


if __name__ == "__main__":
    main()
